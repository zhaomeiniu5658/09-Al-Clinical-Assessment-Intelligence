import re
from pathlib import Path
from typing import Any

from app.models.assessment import ScaleType

HAMD17_ITEMS = [
    "1.抑郁情绪（悲伤、无望、无助、无价值）",
    "2.有罪感",
    "3.自杀",
    "4.入睡困难",
    "5.睡眠不深",
    "6.早醒",
    "7.工作和活动",
    "8.迟滞（指思维和言语缓慢，注意力难以集中，主动性减退）；",
    "9.激越",
    "10.精神性焦虑",
    "11.躯体性焦虑（焦虑的生理症状，如口干、气促、消化不良、腹泻、腹部绞痛、嗳气、心悸、头痛、过度换气、叹气、尿频、出汗）",
    "12.胃肠道症状",
    "13.全身症状",
    "14.性症状（性欲丧失、月经失调）",
    "15.疑病",
    "16.体重减轻",
    "17.自知力",
]


def parse_doctor_test(path: str | Path, scale_type: ScaleType) -> list[dict[str, Any]]:
    """Read the uploaded score sheet into the shape used by the QC result."""
    if scale_type != ScaleType.HAMD:
        raise ValueError("当前仅支持解析 HAMD（HAM-D17）医生打分表")

    file_path = Path(path)
    if not file_path.exists() or not file_path.is_file():
        raise ValueError("医生打分表文件不存在")

    suffix = file_path.suffix.lower()
    if suffix == ".xlsx":
        rows = _read_xlsx(file_path)
    elif suffix == ".xls":
        rows = _read_xls(file_path)
    else:
        raise ValueError("医生打分表仅支持 Excel 格式（.xls 或 .xlsx）")

    return _rows_to_item_results(rows)


def _read_xlsx(path: Path) -> list[tuple[Any, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl 依赖，无法解析 .xlsx 医生打分表") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return [(row[0] if len(row) > 0 else None, row[1] if len(row) > 1 else None) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()


def _read_xls(path: Path) -> list[tuple[Any, Any]]:
    try:
        import xlrd
    except ImportError as exc:
        raise RuntimeError("缺少 xlrd 依赖，无法解析 .xls 医生打分表") from exc

    workbook = xlrd.open_workbook(path, on_demand=True)
    try:
        sheet = workbook.sheet_by_index(0)
        return [
            (sheet.cell_value(row_index, 0) if sheet.ncols > 0 else None, sheet.cell_value(row_index, 1) if sheet.ncols > 1 else None)
            for row_index in range(sheet.nrows)
        ]
    finally:
        workbook.release_resources()


def _rows_to_item_results(rows: list[tuple[Any, Any]]) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    expected_number = 1
    started = False

    for item_value, score_value in rows:
        item_text = _to_text(item_value)
        score = _to_float_or_none(score_value)
        if _is_header(item_text, score_value):
            continue

        explicit_number = _item_number(item_text)
        has_data = explicit_number is not None or score is not None
        if not started and explicit_number != 1:
            continue
        if not has_data:
            if started:
                continue
            continue

        started = True
        item_number = explicit_number or expected_number
        if item_number < 1 or item_number > len(HAMD17_ITEMS):
            break
        if item_number != expected_number:
            item_number = expected_number

        name = item_text or HAMD17_ITEMS[item_number - 1]
        if _item_number(name) is None:
            name = f"{item_number}.{name}"
        results.append(
            {
                "hamd_item": name,
                "doctor_score": score,
                "ai_score": None,
                "ai_scoring_basis": None,
                "difference": None,
            }
        )
        expected_number += 1
        if expected_number > len(HAMD17_ITEMS):
            break

    if not results:
        raise ValueError("医生打分表中未找到 HAM-D17 评分项目")

    while len(results) < len(HAMD17_ITEMS):
        item_number = len(results) + 1
        results.append(
            {
                "hamd_item": HAMD17_ITEMS[item_number - 1],
                "doctor_score": None,
                "ai_score": None,
                "ai_scoring_basis": None,
                "difference": None,
            }
        )
    return results


def _is_header(item_text: str | None, score_value: Any) -> bool:
    text = (item_text or "").strip().lower()
    score_text = _to_text(score_value) or ""
    return "评分员" in text or "医生打分" in text or "项目名称" in text or "评分员" in score_text


def _item_number(value: str | None) -> int | None:
    if not value:
        return None
    match = re.match(r"^\s*(\d{1,2})\s*[.、)\s]", value)
    return int(match.group(1)) if match else None


def _to_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _to_float_or_none(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    return int(number) if number.is_integer() else number
